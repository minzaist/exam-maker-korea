from __future__ import annotations

import html
import io
import re
import zipfile
from pathlib import Path
from typing import Any

import fitz
from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_BREAK, WD_PARAGRAPH_ALIGNMENT
from docx.oxml.ns import qn
from docx.shared import Mm, Pt
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    BaseDocTemplate,
    Frame,
    KeepTogether,
    PageBreak,
    PageTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
)


APP_DIR = Path(__file__).resolve().parent
FONT_PATH = APP_DIR / "assets" / "NanumGothic-Regular.ttf"
FONT_BOLD_PATH = APP_DIR / "assets" / "NanumGothic-Bold.ttf"
_FONT_REGISTERED = False


def _register_fonts() -> tuple[str, str]:
    global _FONT_REGISTERED
    if not _FONT_REGISTERED:
        pdfmetrics.registerFont(TTFont("ExamSans", str(FONT_PATH)))
        bold_path = FONT_BOLD_PATH if FONT_BOLD_PATH.exists() else FONT_PATH
        pdfmetrics.registerFont(TTFont("ExamSansBold", str(bold_path)))
        _FONT_REGISTERED = True
    return "ExamSans", "ExamSansBold"


def _safe(value: Any) -> str:
    return html.escape(str(value or "")).replace("\n", "<br/>")


def _pages(question: dict[str, Any]) -> str:
    values = question.get("pages") or []
    return ", ".join(f"PDF p.{page}" for page in values) or "페이지 확인 필요"


def _page_number(canvas, doc) -> None:
    font, _ = _register_fonts()
    canvas.saveState()
    canvas.setFont(font, 8.5)
    canvas.setFillColor(colors.HexColor("#8B95A1"))
    canvas.drawString(20 * mm, 13 * mm, "시험제작소")
    canvas.drawRightString(A4[0] - 20 * mm, 13 * mm, str(doc.page))
    canvas.restoreState()


def _doc(buffer: io.BytesIO, title: str) -> BaseDocTemplate:
    doc = BaseDocTemplate(
        buffer,
        pagesize=A4,
        rightMargin=19 * mm,
        leftMargin=19 * mm,
        topMargin=20 * mm,
        bottomMargin=20 * mm,
        title=title,
        author="시험제작소",
    )
    frame = Frame(doc.leftMargin, doc.bottomMargin, doc.width, doc.height, id="normal")
    doc.addPageTemplates([PageTemplate(id="main", frames=[frame], onPage=_page_number)])
    return doc


def _styles() -> dict[str, ParagraphStyle]:
    font, bold = _register_fonts()
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "KTitle",
            parent=base["Title"],
            fontName=bold,
            fontSize=22,
            leading=31,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#191F28"),
            spaceAfter=8 * mm,
        ),
        "subtitle": ParagraphStyle(
            "KSubtitle",
            parent=base["Normal"],
            fontName=font,
            fontSize=9.5,
            leading=15,
            alignment=TA_CENTER,
            textColor=colors.HexColor("#6B7684"),
            spaceAfter=8 * mm,
        ),
        "topic": ParagraphStyle(
            "KTopic",
            parent=base["Heading2"],
            fontName=bold,
            fontSize=14,
            leading=20,
            textColor=colors.HexColor("#3182F6"),
            spaceBefore=4 * mm,
            spaceAfter=3 * mm,
        ),
        "question": ParagraphStyle(
            "KQuestion",
            parent=base["Normal"],
            fontName=bold,
            fontSize=11.2,
            leading=18,
            textColor=colors.HexColor("#191F28"),
            spaceAfter=2.5 * mm,
        ),
        "body": ParagraphStyle(
            "KBody",
            parent=base["BodyText"],
            fontName=font,
            fontSize=9.8,
            leading=16,
            textColor=colors.HexColor("#333D4B"),
            spaceAfter=2 * mm,
        ),
        "small": ParagraphStyle(
            "KSmall",
            parent=base["BodyText"],
            fontName=font,
            fontSize=8.3,
            leading=13,
            textColor=colors.HexColor("#8B95A1"),
        ),
        "label": ParagraphStyle(
            "KLabel",
            parent=base["BodyText"],
            fontName=bold,
            fontSize=9.5,
            leading=15,
            textColor=colors.HexColor("#4E5968"),
            spaceBefore=2 * mm,
            spaceAfter=1 * mm,
        ),
    }


def _answer_space(lines: int, width: float) -> Table:
    rows = [[""] for _ in range(max(1, lines))]
    table = Table(rows, colWidths=[width], rowHeights=[6.5 * mm] * len(rows))
    table.setStyle(
        TableStyle(
            [
                ("LINEBELOW", (0, 0), (-1, -1), 0.35, colors.HexColor("#DDE2E8")),
                ("LEFTPADDING", (0, 0), (-1, -1), 0),
                ("RIGHTPADDING", (0, 0), (-1, -1), 0),
            ]
        )
    )
    return table


def make_question_pdf(data: dict[str, Any]) -> bytes:
    buffer = io.BytesIO()
    doc = _doc(buffer, str(data.get("title") or "예상문제집"))
    styles = _styles()
    story: list[Any] = [
        Spacer(1, 12 * mm),
        Paragraph(_safe(data.get("title") or "예상문제집"), styles["title"]),
        Paragraph(_safe(data.get("source_scope") or "첨부 PDF 전체"), styles["subtitle"]),
    ]

    current_topic = None
    for question in data.get("questions", []):
        topic = str(question.get("topic") or "기타")
        if topic != current_topic:
            story.append(Paragraph(_safe(topic), styles["topic"]))
            current_topic = topic
        header = (
            f"{question.get('number')}. {_safe(question.get('question'))} "
            f"<font color='#8B95A1' size='8'>[{_safe(question.get('type'))} · "
            f"{_safe(question.get('difficulty'))}]</font>"
        )
        block: list[Any] = [Paragraph(header, styles["question"])]
        options = question.get("options") or []
        for idx, option in enumerate(options, start=1):
            block.append(Paragraph(f"{idx}) {_safe(option)}", styles["body"]))
        block.append(Paragraph(f"근거 페이지: {_safe(_pages(question))}", styles["small"]))
        lines = int(question.get("answer_lines") or 1)
        if question.get("type") in {"객관식", "OX형"}:
            block.append(Spacer(1, 4 * mm))
            block.append(Paragraph("답: ____________________", styles["body"]))
        elif question.get("type") == "논술형":
            story.extend(block)
            story.append(Spacer(1, 2 * mm))
            story.append(_answer_space(min(lines, 30), doc.width))
            story.append(PageBreak())
            continue
        else:
            block.append(Spacer(1, 2 * mm))
            block.append(_answer_space(lines, doc.width))
        block.append(Spacer(1, 5 * mm))
        story.append(KeepTogether(block))

    doc.build(story)
    return buffer.getvalue()


def make_answer_pdf(data: dict[str, Any]) -> bytes:
    buffer = io.BytesIO()
    doc = _doc(buffer, f"{data.get('title', '예상문제집')} - 정답과 해설")
    styles = _styles()
    story: list[Any] = [
        Spacer(1, 10 * mm),
        Paragraph(f"{_safe(data.get('title') or '예상문제집')}<br/>정답과 해설", styles["title"]),
    ]
    for question in data.get("questions", []):
        answer = question.get("answer") or ""
        if question.get("type") == "객관식" and question.get("answer_index"):
            answer = f"{question.get('answer_index')}번. {answer}"
        block = [
            Paragraph(
                f"{question.get('number')}. {_safe(question.get('question'))}",
                styles["question"],
            ),
            Paragraph("정답 또는 모범답안", styles["label"]),
            Paragraph(_safe(answer), styles["body"]),
            Paragraph("핵심 개념과 주요 근거", styles["label"]),
            Paragraph(_safe(", ".join(question.get("key_concepts") or [])), styles["body"]),
            Paragraph("해설", styles["label"]),
            Paragraph(_safe(question.get("explanation")), styles["body"]),
            Paragraph("풀이 과정", styles["label"]),
            Paragraph(_safe(question.get("reasoning")), styles["body"]),
            Paragraph("주의할 오개념", styles["label"]),
            Paragraph(_safe(question.get("misconceptions")), styles["body"]),
            Paragraph(f"원본 근거: {_safe(_pages(question))}", styles["small"]),
            Spacer(1, 6 * mm),
        ]
        story.append(KeepTogether(block))

    analysis = data.get("analysis") or {}
    story.append(PageBreak())
    story.append(Paragraph("출제 분석 요약", styles["title"]))
    story.append(Paragraph("분석한 시험 범위", styles["label"]))
    story.append(Paragraph(_safe(analysis.get("scope") or data.get("source_scope")), styles["body"]))
    for title, key in (
        ("핵심 출제 개념", "core_concepts"),
        ("중요 학습 주제", "important_topics"),
        ("근거 부족으로 제외한 내용", "excluded_topics"),
    ):
        story.append(Paragraph(title, styles["label"]))
        values = analysis.get(key) or []
        story.append(Paragraph(_safe(" · ".join(map(str, values)) or "없음"), styles["body"]))
    doc.build(story)
    return buffer.getvalue()


def _set_docx_font(document: Document) -> None:
    for style_name in ["Normal", "Title", "Heading 1", "Heading 2"]:
        style = document.styles[style_name]
        style.font.name = "NanumGothic"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "나눔고딕")
    normal = document.styles["Normal"]
    normal.font.size = Pt(10.5)


def _configure_docx(document: Document) -> None:
    section = document.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.left_margin = Mm(20)
    section.right_margin = Mm(20)
    section.top_margin = Mm(20)
    section.bottom_margin = Mm(20)
    _set_docx_font(document)


def make_question_docx(data: dict[str, Any]) -> bytes:
    document = Document()
    _configure_docx(document)
    title = document.add_heading(str(data.get("title") or "예상문제집"), 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    subtitle = document.add_paragraph(str(data.get("source_scope") or "첨부 PDF 전체"))
    subtitle.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    current_topic = None
    for question in data.get("questions", []):
        topic = str(question.get("topic") or "기타")
        if topic != current_topic:
            document.add_heading(topic, level=1)
            current_topic = topic
        paragraph = document.add_paragraph()
        paragraph.add_run(f"{question.get('number')}. {question.get('question')} ").bold = True
        paragraph.add_run(f"[{question.get('type')} · {question.get('difficulty')}]")
        for idx, option in enumerate(question.get("options") or [], 1):
            document.add_paragraph(f"{idx}) {option}")
        document.add_paragraph(f"근거 페이지: {_pages(question)}")
        if question.get("type") in {"객관식", "OX형"}:
            document.add_paragraph("답: ____________________")
        else:
            for _ in range(int(question.get("answer_lines") or 3)):
                document.add_paragraph("________________________________________________________________")
        if question.get("type") == "논술형":
            document.add_page_break()
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def make_answer_docx(data: dict[str, Any]) -> bytes:
    document = Document()
    _configure_docx(document)
    title = document.add_heading(f"{data.get('title', '예상문제집')} - 정답과 해설", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
    for question in data.get("questions", []):
        document.add_heading(f"{question.get('number')}. {question.get('question')}", level=2)
        answer = str(question.get("answer") or "")
        if question.get("type") == "객관식" and question.get("answer_index"):
            answer = f"{question.get('answer_index')}번. {answer}"
        document.add_paragraph(f"정답/모범답안: {answer}")
        document.add_paragraph(f"핵심 개념: {', '.join(question.get('key_concepts') or [])}")
        document.add_paragraph(f"해설: {question.get('explanation') or ''}")
        document.add_paragraph(f"풀이 과정: {question.get('reasoning') or ''}")
        document.add_paragraph(f"주의할 오개념: {question.get('misconceptions') or ''}")
        document.add_paragraph(f"원본 근거: {_pages(question)}")
    document.add_page_break()
    document.add_heading("출제 분석 요약", level=1)
    analysis = data.get("analysis") or {}
    document.add_paragraph(str(analysis.get("scope") or data.get("source_scope") or ""))
    document.add_paragraph("핵심 출제 개념: " + " · ".join(analysis.get("core_concepts") or []))
    document.add_paragraph("중요 학습 주제: " + " · ".join(analysis.get("important_topics") or []))
    document.add_paragraph("근거 부족으로 제외: " + " · ".join(analysis.get("excluded_topics") or []))
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def make_docx_zip(data: dict[str, Any]) -> bytes:
    stream = io.BytesIO()
    with zipfile.ZipFile(stream, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("예상문제집.docx", make_question_docx(data))
        archive.writestr("모범답안_및_해설.docx", make_answer_docx(data))
    return stream.getvalue()


def validate_pdf_bytes(pdf_bytes: bytes, expected_phrase: str = "") -> dict[str, Any]:
    errors: list[str] = []
    try:
        reader = PdfReader(io.BytesIO(pdf_bytes))
        page_count = len(reader.pages)
        extracted = "\n".join((page.extract_text() or "") for page in reader.pages)
        if page_count < 1:
            errors.append("페이지가 없습니다.")
        if expected_phrase and expected_phrase[:12] not in extracted:
            errors.append("제목 텍스트가 PDF에서 확인되지 않습니다.")
        if not extracted.strip():
            errors.append("텍스트 추출 결과가 비어 있습니다.")
    except Exception as exc:
        return {"ok": False, "page_count": 0, "errors": [f"PDF 구조 오류: {exc}"]}

    try:
        document = fitz.open(stream=pdf_bytes, filetype="pdf")
        for index, page in enumerate(document):
            rect = page.rect
            if abs(rect.width - 595.28) > 4 or abs(rect.height - 841.89) > 4:
                errors.append(f"{index + 1}페이지가 A4 규격이 아닙니다.")
            pixmap = page.get_pixmap(matrix=fitz.Matrix(0.5, 0.5), alpha=False)
            if len(set(pixmap.samples)) < 2:
                errors.append(f"{index + 1}페이지 렌더링이 비어 있습니다.")
        document.close()
    except Exception as exc:
        errors.append(f"PDF 렌더링 검사 오류: {exc}")
    return {"ok": not errors, "page_count": page_count, "errors": errors}


def natural_language_output(data: dict[str, Any]) -> str:
    lines = [str(data.get("title") or "예상문제집"), "", str(data.get("source_scope") or "")]
    for question in data.get("questions", []):
        lines.extend(
            [
                "",
                f"[{question.get('topic')}]",
                f"{question.get('number')}. {question.get('question')} ({question.get('type')} · {question.get('difficulty')})",
            ]
        )
        for idx, option in enumerate(question.get("options") or [], 1):
            lines.append(f"  {idx}) {option}")
        lines.append(f"  근거: {_pages(question)}")
    lines.extend(["", "정답 및 해설"])
    for question in data.get("questions", []):
        answer = str(question.get("answer") or "")
        if question.get("type") == "객관식" and question.get("answer_index"):
            answer = f"{question.get('answer_index')}번. {answer}"
        lines.extend(
            [
                "",
                f"{question.get('number')}. {answer}",
                f"해설: {question.get('explanation') or ''}",
                f"풀이: {question.get('reasoning') or ''}",
                f"주의: {question.get('misconceptions') or ''}",
                f"근거: {_pages(question)}",
            ]
        )
    return "\n".join(lines)

